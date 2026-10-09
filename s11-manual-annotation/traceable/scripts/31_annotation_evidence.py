#!/opt/homebrew/bin/python3.13
"""Annotation evidence: locate every value cell of every manual annotation assay in the protocol markdown the LLM reads.

For each annotation value cell (Characteristic / Parameter / Factor / Comment / Date column, every sheet, as in
28_score_extraction.py) the cell is normalised exactly like the scorer does (28's atoms(): dates -> ISO,
decimal comma -> point, case, punctuation), and its token sequence is searched in the tokenised source text.
The source is tokenised position-aware with the same rules, after masking HTML tags (<td>, <mark>, <u> ...) so
markup and line breaks never split a value. A number glued to a known unit ("48h", "37°C", "180µl") is split into
number + unit, so the annotation value 48 is found in "48h"; identifiers are not split ("7mu", "R05", "BAS33"). A number written in digit groups in the source ("97 613 732") also
matches the ungrouped annotation value ("97613732").

Every occurrence is kept (ExtractBench's "evidence list": any listed location is accepted). A cell is
  unique    exactly one occurrence  -> its location is certain ("verified" annotation evidence)
  multiple  several occurrences     -> any of them is accepted, but the right one is not known
  none      no verbatim occurrence  -> derived, normalised or reworded by the curator (e.g. 26.5 h for 26h30min)
Hash cells are skipped (as in 28).

Context resolution (B): a "multiple" cell is resolved to one occurrence from the cell's own row, the way a curator reads
the notebook (run ID, date, strain / sample label, protocol step). The row's identifying tokens are the words and dates of
its other cells (names split at _ - # and spaces; numbers excluded, they are the ambiguous part; the sheet name), minus tokens that half
or more of the rows of the sheet share (they identify the table, not the row). Each
occurrence is scored by how close those tokens stand to it in the source: same line / same HTML table row 1.0, within one
line 0.8, three lines 0.5, fifteen lines 0.15, plus 0.3 when the token lies in the same section (between headings or
lone date lines); each token weighted 1/sqrt(its number of occurrences in the source) so a rare run ID outweighs a common word. The occurrence with the strictly
highest score is taken when it leads the second best by at least MARGIN; else the cell stays "multiple". Status "context".
All occurrences stay listed; `resolved` indexes the chosen one. Context-resolved cells are weaker evidence than "unique"
ones (a heuristic, not a certainty): summaries report them separately.

Usage: 31_annotation_evidence.py [phage|sdy1|biolector034|ux511|pux782|all]   (default all)
  -> work/eval[_<case>]/annotation_evidence.<annotation>.json (source file + sha256, per-cell occurrences, summary)
     and work/eval/annotation_evidence_summary.md
"""
import bisect
import hashlib
import importlib.util
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent                  # root of the section; the run folders (work/eval*/) are not part of it
_spec = importlib.util.spec_from_file_location("score", ROOT / "scripts" / "28_score_extraction.py")
S = importlib.util.module_from_spec(_spec)
_argv, sys.argv = sys.argv, sys.argv[:1]
_spec.loader.exec_module(S)
sys.argv = _argv

TAG = re.compile(r"<[^<>]{0,300}>")
TOKEN = re.compile(r"[a-z0-9äöüßµμ°%]+(?:[.,]\d+)?", re.I)
GLUED = re.compile(r"(\d+(?:[.,]\d+)?)([a-zµμ°%]+)")
UNIT_WORDS = {u for u in S.UNITS if "/" not in u and " " not in u} | {"x", "xg", "rpm", "°c", "%", "kb", "bp", "nt"}


def source_path(case):
    return next((S.CASES[case]["eval"] / "input").glob("*.md"))


def tokenize_source(text):
    """-> list of (normalised token, start, end) over the source; dates as one ISO token, a dash between two
    numbers as "-", HTML tags masked (same length, so offsets stay valid)."""
    masked = TAG.sub(lambda m: " " * len(m.group(0)), text)
    spans = []                                   # (start, end, normalised) of dates first, then plain tokens
    taken = [False] * len(masked)
    for m in S.DATE.finditer(masked):
        if m.group(1):
            y, mo, d = m.group(1), m.group(2), m.group(3)
        else:
            d, mo, y = m.group(4), m.group(5), m.group(6)
            y = y if len(y) == 4 else "20" + y
        spans.append((m.start(), m.end(), f"{y}-{int(mo):02d}-{int(d):02d}"))
        for k in range(m.start(), m.end()):
            taken[k] = True
    for m in TOKEN.finditer(masked):
        if any(taken[m.start():m.end()]):
            continue
        tok = m.group(0).lower()
        glued = GLUED.fullmatch(tok)
        if glued and glued.group(2) in UNIT_WORDS:  # "48h", "37°c", "180µl": the number and the unit as two tokens
            cut = m.start() + len(glued.group(1))
            spans.append((m.start(), cut, S.canonical_number(glued.group(1))))
            spans.append((cut, m.end(), glued.group(2)))
            continue
        norm = S.canonical_number(tok) if re.fullmatch(r"\d+(?:[.,]\d+)?", tok) else tok
        spans.append((m.start(), m.end(), norm))
    spans.sort()
    out = []
    for i, (a, b, t) in enumerate(spans):        # "15-16": a dash between two numbers is its own token
        if out and re.fullmatch(r"\d+(?:\.\d+)?", t) and re.fullmatch(r"\d+(?:\.\d+)?", out[-1][0]) \
                and re.fullmatch(r"\s*[-–]\s*", masked[out[-1][2]:a]):
            out.append(("-", out[-1][2], a))
        out.append((t, a, b))
    return out, masked


def find(tokens, masked, cell_tokens):
    """Character spans of every occurrence of a normalised cell (token sequence) in the source."""
    seq = []
    for tok in cell_tokens.split():               # split glued number + unit exactly as in the source
        glued = GLUED.fullmatch(tok)
        seq += [S.canonical_number(glued.group(1)), glued.group(2)] if glued and glued.group(2) in UNIT_WORDS else [tok]
    n = len(seq)
    hits = [(tokens[i][1], tokens[i + n - 1][2]) for i in range(len(tokens) - n + 1)
            if all(tokens[i + k][0] == seq[k] for k in range(n))]
    if not hits and n == 1 and re.fullmatch(r"\d{6,}", seq[0]):   # "97613732" written as "97 613 732"
        grouped = r"(?<!\d)" + r"\s?".join(seq[0]) + r"(?!\d)"
        hits = [(m.start(), m.end()) for m in re.finditer(grouped, masked)]
    return hits


SKIP_COLUMN = re.compile(r"^(Term Source REF|Term Accession Number|Unit|Protocol Type|Protocol Uri|Protocol REF|Output \[Data\]|"
                         r"Comment \[(sha256|original file name)\])", re.I)
MARGIN = 0.25                                      # lead over the second best occurrence needed to resolve a cell
PROXIMITY = ((0, 1.0), (1, 0.8), (3, 0.5))         # (distance in text units, weight); then section 0.3, 15 units 0.15


def row_context(sheet, header, row, skip):
    """Identifying tokens of one annotation row: words and dates of its cells except the cell itself (`skip` = column
    indexes to leave out) and structural columns; the sheet name; no plain numbers."""
    parts = [sheet.replace("_", " ")]
    for j, v in enumerate(row):
        if j in skip or v in (None, "") or SKIP_COLUMN.match(header[j]):
            continue
        parts.append(str(v))
    toks, _ = tokenize_source(re.sub(r"[_#]+|(?<!\d)-+|-+(?!\d)", " ", " ; ".join(parts)))
    return sorted({t for t, _, _ in toks if len(t) > 1 and not re.fullmatch(r"\d+(?:\.\d+)?", t)})


def annotation_cells(xlsx):
    """-> list of annotation value cells: sheet, row (1-based data row), column header, value, unit, normalised."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    cells = []
    for name in (s for s in wb.sheetnames if s != "isa_assay"):
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        header = [str(h or "").strip() for h in rows[0]]
        for j, h in enumerate(header):
            if not S.VALUE_HEADER.match(h):
                continue
            has_unit = j + 1 < len(header) and header[j + 1] == "Unit"
            for i, r in enumerate(rows[1:], start=1):
                value = r[j] if j < len(r) and r[j] is not None else ""
                unit = (r[j + 1] or "") if has_unit and j + 1 < len(r) else ""
                for kind, fact, lint_form in S.atoms(value, unit):
                    cells.append({"sheet": name, "row": i, "column": h, "value": str(value), "unit": str(unit),
                                  "kind": kind, "normalised": fact, "lint": lint_form,
                                  "context": row_context(name, header, r, {j, j + 1} if has_unit else {j})})
    return row_specific(cells)


def row_specific(cells):
    """Drop context tokens that half or more of the rows of the cell's sheet share (sheet name, medium, "assay"): they
    identify the table, not the row, and would pull every row towards the same occurrence."""
    rows, freq = {}, {}
    for c in cells:
        rows.setdefault(c["sheet"], set()).add(c["row"])
    seen = set()
    for c in cells:
        for t in c["context"]:
            if (c["sheet"], c["row"], t) not in seen:
                seen.add((c["sheet"], c["row"], t))
                freq[c["sheet"], t] = freq.get((c["sheet"], t), 0) + 1
    for c in cells:
        n = len(rows[c["sheet"]])
        c["context"] = [t for t in c["context"] if n < 4 or freq[c["sheet"], t] < 0.5 * n]
    return cells


class Layout:
    """Text units of the source (a line, or a whole HTML <tr>...</tr>) and sections (between headings / lone date
    lines), so that distances are measured the way a reader sees the notebook."""
    LONE_DATE = re.compile(r"^\W*(?:\d{4}-\d{2}-\d{2}|\d{1,2}[./]\d{1,2}[./](?:\d{2}|\d{4}))\W*$")

    def __init__(self, text):
        self.starts = [0]
        for m in re.finditer(r"\n", text):
            self.starts.append(m.end())
        self.unit_of_line = list(range(len(self.starts)))
        for m in re.finditer(r"<tr\b.*?</tr>", text, re.S | re.I):          # a table row spanning lines = one unit
            a, b = self.line(m.start()), self.line(m.end())
            for k in range(a, b + 1):
                self.unit_of_line[k] = self.unit_of_line[a]
        lines = text.split("\n")
        self.section_of_line, sec = [], 0
        for ln in lines:
            if ln.lstrip().startswith("#") or self.LONE_DATE.match(ln):
                sec += 1
            self.section_of_line.append(sec)

    def line(self, pos):
        return bisect.bisect_right(self.starts, pos) - 1

    def unit(self, pos):
        return self.unit_of_line[self.line(pos)]

    def section(self, pos):
        return self.section_of_line[self.line(pos)]


UNIT_DIM = {**{u: d for u, (d, _) in S.UNITS.items()}, "percent": "pct", "%": "pct"}
UNIT_AFTER = re.compile(rf"\s{{0,2}}({S.UNIT_RX}|%)(?![a-zA-Z])", re.I)


def unit_fit(cell, masked, end):
    """+1 when the number at `end` is followed by a unit of the annotation cell's dimension, -inf when followed by a unit of
    another dimension (5 uL is not the 5 mg/ml), 0 when either side has no known unit."""
    annotation = UNIT_DIM.get(cell["unit"].strip().lower()) if cell["unit"] else None
    m = UNIT_AFTER.match(masked, end) if annotation else None
    if not m:
        return 0.0
    return 1.0 if UNIT_DIM[m.group(1).lower()] == annotation else -math.inf


def resolve_context(cell, hits, layout, positions, masked):
    """Index of the occurrence (in `hits`) the cell's row context points to, with (best, second) scores, or None.
    `positions`: context token -> sorted character offsets in the source. The nearest other occurrence of a token on
    each side of the cell decides its distance (units and sections are monotonic along the text)."""
    if len(hits) < 2 or not cell["context"]:
        return None, 0.0, 0.0
    scores = []
    for a, b in hits:
        ua, sa = layout.unit(a), layout.section(a)
        total = 0.0
        for t in cell["context"]:
            pos = positions.get(t)
            if not pos:
                continue
            i, j = bisect.bisect_left(pos, a), bisect.bisect_left(pos, b)
            near = [p for p in (pos[i - 1] if i else None, pos[j] if j < len(pos) else None) if p is not None]
            if not near:
                continue
            d = min(abs(layout.unit(p) - ua) for p in near)
            w = next((w for lim, w in PROXIMITY if d <= lim), 0.15 if d <= 15 else 0.0)
            if any(layout.section(p) == sa for p in near):
                w += 0.3                                  # same section as the token: it belongs to this run / date
            total += w / math.sqrt(len(pos))
        scores.append(total)
    fits = [unit_fit(cell, masked, b) for _, b in hits]
    keep = [f > 0 for f in fits] if any(f > 0 for f in fits) else [f == 0 for f in fits]   # unit match is decisive
    context_top = max(scores)
    raw = list(scores)
    scores = [sc if k else -math.inf for sc, k in zip(scores, keep)]
    order = sorted(range(len(hits)), key=lambda i: -scores[i])
    best, second = scores[order[0]], scores[order[1]]
    if second == -math.inf:                      # one unit-compatible candidate left: only if the context agrees
        return (order[0] if raw[order[0]] >= context_top - 1e-9 and best > 0 else None), best, second
    return (order[0] if best > 0 and best - second >= MARGIN else None), best, second


def line_of(text, pos):
    return text.count("\n", 0, pos) + 1


def build(case):
    src_file = source_path(case)
    text = src_file.read_text()
    tokens, masked = tokenize_source(text)
    layout = Layout(text)
    positions = {}
    for t, a, _ in tokens:
        positions.setdefault(t, []).append(a)
    summaries, by_annotation = [], {}
    for annotation, xlsx in S.CASES[case]["annotations"].items():
        cells = annotation_cells(xlsx)
        cache = {}
        for c in cells:
            if c["normalised"] not in cache:
                cache[c["normalised"]] = find(tokens, masked, c["normalised"])
            hits = cache[c["normalised"]]
            c["occurrences"] = [{"char_span": [a, b], "line": line_of(text, a), "text": text[a:b]} for a, b in hits]
            c["status"] = "unique" if len(hits) == 1 else "multiple" if hits else "none"
            if c["status"] == "multiple":
                k, best, second = resolve_context(c, hits, layout, positions, masked)
                c["context_scores"] = [round(best, 3), round(second, 3)]
                if k is not None:
                    c["status"], c["resolved"] = "context", k
                    sec = layout.section(hits[k][0])
                    c["same_section"] = [i for i, (a, _) in enumerate(hits) if layout.section(a) == sec]
        count = {s: sum(1 for c in cells if c["status"] == s) for s in ("unique", "context", "multiple", "none")}
        slug = re.sub(r"\W+", "_", annotation.lower()).strip("_")
        out = S.CASES[case]["eval"] / f"annotation_evidence.{slug}.json"
        (_, same, tested, by_unit), = shuffle_control(case, {annotation: cells})
        row_specific_ok = tested == 0 or same / tested < 0.5
        print(f"   shuffle control {case} / {annotation}: with another sheet's context {same}/{tested} resolved cells keep their "
              f"occurrence ({by_unit} more are decided by their unit alone) -> context resolution "
              f"{'is row-specific' if row_specific_ok else 'is NOT row-specific here (single-section source): not used as verified'}")
        out.write_text(json.dumps({
            "case": case, "annotation": annotation, "annotation_xlsx": str(xlsx.relative_to(ROOT)),
            "source": str(src_file.relative_to(ROOT)), "source_sha256": hashlib.sha256(text.encode()).hexdigest(),
            "summary": {"cells": len(cells), **count, "context_row_specific": row_specific_ok,
                        "shuffle_same_of_tested": [same, tested]}, "cells": cells}, indent=1, ensure_ascii=False))
        summaries.append((case, annotation, len(cells), count, out))
        by_annotation[annotation] = cells
        print(f"{case} / {annotation}: {len(cells)} cells - unique {count['unique']}, context {count['context']}, "
              f"multiple {count['multiple']}, none {count['none']} -> {out.relative_to(ROOT)}")
    return summaries


def control():
    """The locator must find a value across markup, a line break and a decimal comma, and not inside a longer number."""
    text = "OD <mark>0,653</mark> at 15\nmg/mL; code 97 613 732; not 10.6532; grow 48h at 37°C; strain 7mu"
    tokens, masked = tokenize_source(text)
    assert [text[a:b] for a, b in find(tokens, masked, "0.653")] == ["0,653"]
    assert [text[a:b] for a, b in find(tokens, masked, "15 mg ml")] == ["15\nmg/mL"]
    assert [text[a:b] for a, b in find(tokens, masked, "97613732")] == ["97 613 732"]
    assert find(tokens, masked, "0.6532") == [] and len(find(tokens, masked, "10.6532")) == 1
    assert [text[a:b] for a, b in find(tokens, masked, "48")] == ["48"] and len(find(tokens, masked, "37")) == 1
    assert find(tokens, masked, "7") == [] and len(find(tokens, masked, "7mu")) == 1
    assert len(find(tokens, masked, "48h")) == 1               # a glued annotation value matches the glued source
    print("control OK: markup, line break, decimal comma and digit groups are matched; partial numbers are not")


def prepared(text):
    tokens, masked = tokenize_source(text)
    positions = {}
    for t, a, _ in tokens:
        positions.setdefault(t, []).append(a)
    return tokens, masked, Layout(text), positions


def context_control():
    """B3. Context resolution must pick the occurrence the row points to, and must stay unresolved when the context is
    absent, tied, or contradicted by the unit."""
    text = ("# Run R04\nincubate at 37 C for 3 h; add 5 uL of stock\n# Run R05\nincubate at 37 C for 3 h; "
            "stock is 5 mg/ml\n")
    tokens, masked, layout, positions = prepared(text)
    hits = find(tokens, masked, "37")
    assert len(hits) == 2
    cell = lambda ctx, unit="": {"context": ctx, "unit": unit}
    assert resolve_context(cell(["r04"]), hits, layout, positions, masked)[0] == 0
    assert resolve_context(cell(["r05"]), hits, layout, positions, masked)[0] == 1            # wrong context -> wrong place
    assert resolve_context(cell(["zzz"]), hits, layout, positions, masked)[0] is None         # context absent
    t3, m3, l3, p3 = prepared("# Run A\nincubate at 37 C\n# Run A\nincubate at 37 C\n")
    assert resolve_context(cell(["a"]), find(t3, m3, "37"), l3, p3, m3)[0] is None            # identical contexts: a tie
    assert resolve_context(cell([]), hits, layout, positions, masked)[0] is None
    five = find(tokens, masked, "5")
    assert len(five) == 2
    assert resolve_context(cell(["r04"], "milligram per milliliter"), five, layout, positions, masked)[0] is None  # unit vs context disagree
    assert resolve_context(cell([], "microliter"), five, layout, positions, masked)[0] is None                   # no context at all
    assert resolve_context(cell(["r05"], "milligram per milliliter"), five, layout, positions, masked)[0] == 1
    assert resolve_context(cell(["r04"], "microliter"), five, layout, positions, masked)[0] == 0
    # unit decides between two lines the context cannot separate
    text2 = "# Run R04\nmix 5 uL and 5 mg/ml\n"
    t2, m2, l2, p2 = prepared(text2)
    h2 = find(t2, m2, "5")
    assert resolve_context(cell(["r04"], "microliter"), h2, l2, p2, m2)[0] == 0
    assert resolve_context(cell(["r04"], "milligram per milliliter"), h2, l2, p2, m2)[0] == 1
    print("context control OK: right run picked, wrong run picked for wrong context, absent / tied / unit-contradicted context stays multiple")


def shuffle_control(case, cells_by_annotation):
    """On real data: the same resolver with another sheet's row context must mostly choose different occurrences.
    Cells decided by their unit alone (exactly one occurrence followed by the right unit) are reported apart: they do
    not depend on the context by design."""
    text = source_path(case).read_text()
    tokens, masked, layout, positions = prepared(text)
    out = []
    for annotation, cells in cells_by_annotation.items():
        other = [c for c in cells if c["context"]]
        same = tested = by_unit = 0
        for i, c in enumerate(c for c in cells if c["status"] == "context"):
            hits = [tuple(o["char_span"]) for o in c["occurrences"]]
            if sum(unit_fit(c, masked, b) > 0 for _, b in hits) == 1:
                by_unit += 1
                continue
            donor = next((d for d in other[(i * 37) % len(other):] + other if d["sheet"] != c["sheet"]), None)
            if donor is None:
                continue
            k = resolve_context(dict(c, context=donor["context"]), hits, layout, positions, masked)[0]
            tested += 1
            same += k == c["resolved"]
        out.append((annotation, same, tested, by_unit))
    return out


def main():
    control()
    context_control()
    arg = sys.argv[1] if len(sys.argv) > 1 else "all"
    rows = []
    for case in (list(S.CASES) if arg == "all" else [arg]):
        rows += build(case)
    if arg == "all":
        md = ["# How traceable are the manual annotations by lookup?", "",
              "Generated by `scripts/31_annotation_evidence.py`. Each annotation value cell searched in the markdown the LLM reads. "
              "**unique** = exactly one location (certain), **context-resolved** = several candidate locations, one chosen from the row's own context (run / date / sample / step; a heuristic, see the docstring), **multiple** = several candidate locations left, **none** = not "
              "written verbatim in the source (derived, normalised or reworded by the curator).", "",
              "| Entry | Annotation | Value cells | Unique | Context-resolved | Multiple | None |", "|---|---|---|---|---|---|---|"]
        for case, annotation, n, c, _ in rows:
            md.append(f"| {case}{' (rendered-page HTML input)' if case == 'pux782' else ''} | {annotation} | {n} | {c['unique']} ({c['unique'] / n:.0%}) | {c['context']} ({c['context'] / n:.0%}) | {c['multiple']} "
                      f"({c['multiple'] / n:.0%}) | {c['none']} ({c['none'] / n:.0%}) |")
        (ROOT / "work/eval/annotation_evidence_summary.md").write_text("\n".join(md) + "\n")
        print("\n".join(md))


if __name__ == "__main__":
    main()
